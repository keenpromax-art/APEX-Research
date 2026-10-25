from __future__ import annotations

import io
import re
import unicodedata
from datetime import datetime, timezone
from typing import Any
from xml.sax.saxutils import escape

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/api/export", tags=["export"])

# --- Small-report PDF helpers -------------------------------------------------
# Everything below is dedicated to the short APEX Terminal research note
# (POST /api/export/pdf). The large institutional report is rendered by
# services/pdf_report.py and must not share or depend on this code.

_INK = "#0F172A"
_BODY = "#1E293B"
_MUTED = "#64748B"
_ACCENT = "#1D4ED8"
_RULE = "#E2E8F0"
_PANEL = "#F8FAFC"
_WARN_BG = "#FEF3C7"
_WARN_BORDER = "#FCD34D"
_WARN_TEXT = "#92400E"
_NEG = "#B91C1C"

_SECTION_ORDER = [
    "executive_summary",
    "investment_thesis",
    "business_overview",
    "financial_performance",
    "strengths",
    "risks",
    "catalysts",
    "valuation_view",
    "forecast_assumptions",
    "monitoring",
    "conclusion",
    "sources_limitations",
    "disclaimer",
]

# Characters Helvetica (WinAnsi) cannot render, folded to safe ASCII equivalents.
# En dash (U+2013) and em dash (U+2014) are part of cp1252 and must survive.
_FOLD = {code: "-" for code in (0x2010, 0x2011, 0x2012, 0x2015)}
_FOLD.update(
    {
        0x2212: "-",  # minus sign
        0x2248: "~",  # almost equal
        0x2264: "<=",
        0x2265: ">=",
        0x2192: "->",
        0x2026: "...",
        0x20B9: "INR ",
        0x00A0: " ",
        0x2044: "/",
    }
)


def _sanitize(value: Any) -> str:
    """Fold unicode variants and drop glyphs Helvetica cannot draw (no tofu boxes)."""
    text = unicodedata.normalize("NFKC", str(value)).replace("\u00ad", "")
    out: list[str] = []
    for char in text:
        replacement = _FOLD.get(ord(char))
        if replacement is not None:
            out.append(replacement)
            continue
        try:
            char.encode("cp1252")
            out.append(char)
        except UnicodeEncodeError:
            folded = unicodedata.normalize("NFKD", char).encode("ascii", "ignore").decode("ascii")
            out.append(folded)
    return "".join(out)


def _f(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if result != result or result in (float("inf"), float("-inf")):
        return None
    return result


def _fmt_compact(value: Any) -> str:
    number = _f(value)
    if number is None:
        return "—"
    sign = "-" if number < 0 else ""
    amount = abs(number)
    if amount >= 1e9:
        return f"{sign}{amount / 1e9:,.2f} bn"
    if amount >= 1e6:
        return f"{sign}{amount / 1e6:,.2f} m"
    if amount >= 1:
        return f"{sign}{amount:,.2f}"
    return f"{sign}{amount:,.4f}"


def _fmt_money(value: Any, currency: str) -> str:
    number = _f(value)
    if number is None:
        return "—"
    sign = "-" if number < 0 else ""
    code = (currency or "").upper()
    prefix = f"{code} " if code else ""
    return f"{sign}{prefix}{_fmt_compact(abs(number))}"


def _fmt_pct(value: Any, digits: int = 1, signed: bool = False) -> str | None:
    number = _f(value)
    if number is None:
        return None
    return f"{number * 100:+.{digits}f}%" if signed else f"{number * 100:.{digits}f}%"


def _fmt_x(value: Any) -> str | None:
    number = _f(value)
    if number is None:
        return None
    return f"{number:,.2f}x"


def _fy_year(period_end: Any) -> int | None:
    match = re.match(r"(\d{4})", str(period_end or ""))
    return int(match.group(1)) if match else None


_ACRONYMS = {
    "wacc": "WACC",
    "ev": "EV",
    "ebit": "EBIT",
    "ebitda": "EBITDA",
    "eps": "EPS",
    "fcf": "FCF",
    "ocf": "OCF",
    "dcf": "DCF",
    "roe": "ROE",
    "roa": "ROA",
    "roic": "ROIC",
}


def _humanize(key: str) -> str:
    return " ".join(_ACRONYMS.get(part, part.title()) for part in key.split("_"))


_DCF_PERCENT_KEYS = {
    "wacc",
    "terminal_growth",
    "terminal_growth_rate",
    "tax_rate",
    "risk_free",
    "risk_free_rate",
    "market_premium",
    "cost_of_equity",
    "cost_of_debt",
    "growth_rate",
    "implied_growth",
    "discount_rate",
}


def _fmt_dcf_value(key: str, value: Any, currency: str) -> str:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    number = _f(value)
    if number is None:
        return _sanitize(value)
    lowered = key.lower()
    if lowered in _DCF_PERCENT_KEYS:
        return f"{number * 100:.1f}%"
    if "multiple" in lowered:
        return f"{number:,.1f}x"
    if "per_share" in lowered or lowered in {"fair_value", "price_target", "target_price"}:
        prefix = f"{currency.upper()} " if currency else ""
        return f"{prefix}{number:,.2f}"
    if abs(number) >= 1e6:
        return _fmt_compact(number)
    if isinstance(value, int) or float(number).is_integer() and abs(number) >= 1000:
        return f"{number:,.0f}"
    return f"{number:,.2f}"


def _kpi_rows(body: dict[str, Any]) -> list[tuple[str, str, str]]:
    company = body.get("company") or {}
    price = body.get("price_summary") or {}
    actuals = body.get("actuals") or {}
    currency = str(company.get("currency") or "")
    fy = _fy_year(company.get("last_data_date"))
    fy_label = f"FY{fy}" if fy else "latest year"

    close = _f(price.get("last_close"))
    change_pct = _f(price.get("change_pct"))
    eps = _f(actuals.get("eps_diluted"))
    bvps = _f(actuals.get("book_value_per_share"))

    rows: list[tuple[str, str, str]] = []
    if close is not None:
        sub = f"{change_pct * 100:+.1f}% vs prev. close" if change_pct is not None else fy_label
        rows.append(("Price", _fmt_money(close, currency), sub))
    market_cap = _f(company.get("market_cap"))
    if market_cap is not None:
        rows.append(("Market cap", _fmt_money(market_cap, currency), currency or "as reported"))
    if close is not None and eps is not None and eps > 0:
        rows.append(("P/E (trailing)", f"{close / eps:,.2f}x", f"price / {fy_label} EPS"))
    if close is not None and bvps is not None and bvps > 0:
        rows.append(("Price / book", f"{close / bvps:,.2f}x", "price / book value"))
    growth = _fmt_pct(actuals.get("revenue_growth"), signed=True)
    if growth:
        rows.append(("Revenue growth", growth, f"{fy_label} YoY"))
    operating_margin = _fmt_pct(actuals.get("operating_margin"))
    if operating_margin:
        rows.append(("Operating margin", operating_margin, fy_label))
    leverage = _fmt_x(actuals.get("net_debt_ebitda"))
    if leverage:
        rows.append(("Net debt / EBITDA", leverage, "leverage"))
    fcf = _f(actuals.get("fcf"))
    if fcf is not None:
        rows.append(("Free cash flow", _fmt_money(fcf, currency), fy_label))
    return rows


_FORECAST_LINES = [("Revenue", "revenue"), ("EBITDA", "ebitda"), ("Net income", "net_income"), ("Free cash flow (FCF)", "fcf")]


def _recommendation_rows(body: dict[str, Any]) -> list[tuple[str, str, str]]:
    rating = body.get("rating") if isinstance(body.get("rating"), dict) else {}
    dcf = body.get("dcf") if isinstance(body.get("dcf"), dict) else {}
    price = body.get("price_summary") if isinstance(body.get("price_summary"), dict) else {}
    currency = str((body.get("company") or {}).get("currency") or "")
    target = _f(rating.get("target_price"))
    if target is None:
        target = _f(dcf.get("value_per_share"))
    current = _f(price.get("last_close") or price.get("current"))
    upside = _f(rating.get("upside"))
    if upside is None and target is not None and current:
        upside = target / current - 1
    code = str(rating.get("code") or ("NOT RATED" if target is None else ""))
    rows: list[tuple[str, str, str]] = []
    if code:
        rows.append(("Rating", _sanitize(code), str(rating.get("horizon") or "12M") + " horizon"))
    if target is not None:
        rows.append(("Model target (E)", _fmt_money(target, currency), "APEX deterministic DCF"))
    if current is not None:
        date = _sanitize(str(price.get("last_date") or price.get("price_date") or ""))
        rows.append(("Last close", _fmt_money(current, currency), f"as of {date}" if date else "last available close"))
    if upside is not None:
        rows.append(("Implied upside", _fmt_pct(upside, signed=True) or "—", "target ÷ last close − 1"))
    if dcf.get("wacc") is not None:
        rows.append(("Model WACC", _fmt_pct(dcf.get("wacc")) or "—", "DCF discount rate"))
    return rows


def _price_chart_png(price_history: list[dict[str, Any]], target: Any, currency: str) -> bytes | None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None
    series = [
        {"date": str(point.get("date") or ""), "close": _f(point.get("close"))}
        for point in price_history[-260:]
    ]
    series = [point for point in series if point["close"] is not None]
    if len(series) < 12:
        return None
    closes = [point["close"] for point in series]
    target_value = _f(target)
    low, high = min(closes), max(closes)
    if target_value is not None:
        low, high = min(low, target_value), max(high, target_value)
    pad = (high - low) or abs(high) or 1.0
    low, high = low - pad * 0.14, high + pad * 0.14
    fig, ax = plt.subplots(figsize=(7.0, 2.0))
    ax.fill_between(range(len(closes)), closes, low, color="#93B4F5", alpha=0.30, linewidth=0)
    ax.plot(range(len(closes)), closes, color="#1D4ED8", linewidth=1.2)
    tick_indices = sorted(set([round(i * (len(series) - 1) / 4) for i in range(5)]))
    ax.set_xticks(tick_indices)
    ax.set_xticklabels([series[i]["date"][:7] for i in tick_indices], fontsize=6.0)
    ax.set_ylim(low, high)
    ax.set_ylabel(currency.upper() or "price", fontsize=6.5, color="#64748B")
    if target_value is not None:
        ax.axhline(target_value, color="#B45309", linewidth=0.9, linestyle=(0, (4, 3)))
        ax.annotate(
            f"model target {_fmt_money(target_value, currency)}",
            (0.02, target_value),
            xycoords=("axes fraction", "data"),
            fontsize=6.0,
            color="#B45309",
            ha="left",
            va="bottom",
        )
    ax.annotate(
        _fmt_money(closes[-1], currency),
        (len(closes) - 1, closes[-1]),
        xytext=(-2, 4),
        textcoords="offset points",
        fontsize=6.5,
        color="#1D4ED8",
        fontweight="bold",
        ha="right",
        va="bottom",
    )
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#94A3B8")
    ax.tick_params(colors="#64748B", labelsize=6.2, length=0)
    ax.grid(axis="y", color="#E2E8F0", linewidth=0.5)
    ax.set_axisbelow(True)
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer.read()


def _has_actuals(actuals: dict[str, Any]) -> bool:
    return any(_f(actuals.get(key)) is not None for _label, key in _FORECAST_LINES)


def _filename(body: dict[str, Any], extension: str) -> str:
    ticker = str(body.get("ticker") or "model").upper()
    safe = re.sub(r"[^A-Z0-9._-]+", "_", ticker).strip("._") or "model"
    return f"apex_{safe}_{extension}"


def _scalar(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


@router.post("/xlsx")
def export_xlsx(body: dict[str, Any]):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"Excel export unavailable: {error}")

    workbook = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    bold = Font(bold=True)

    def style_header(sheet, row: int = 1) -> None:
        for cell in sheet[row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    def autosize(sheet) -> None:
        for column in sheet.columns:
            length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(length + 2, 10), 48)

    forecast = body.get("forecast") or {}
    periods = forecast.get("periods") or []
    sheet = workbook.active
    sheet.title = "Forecast"
    if periods:
        lines = [
            "revenue",
            "gross_profit",
            "ebitda",
            "ebit",
            "net_income",
            "ocf",
            "fcf",
            "capex",
            "cash",
            "debt",
            "equity",
            "eps",
            "balance_check",
        ]
        sheet.append(["Line item"] + [f"Y{period.get('year', index + 1)}" for index, period in enumerate(periods)])
        for line in lines:
            sheet.append([line.replace("_", " ").title()] + [period.get(line) for period in periods])
        style_header(sheet)
        sheet.freeze_panes = "B2"
        autosize(sheet)
    else:
        sheet.append(["No forecast supplied"])
        autosize(sheet)

    assumptions = forecast.get("assumptions") or body.get("assumptions") or {}
    if assumptions:
        sheet = workbook.create_sheet("Assumptions")
        sheet.append(["Assumption", "Value"])
        for key, value in assumptions.items():
            sheet.append([key.replace("_", " ").title(), _scalar(value)])
        style_header(sheet)
        autosize(sheet)

    dcf_result = body.get("dcf") or {}
    if dcf_result:
        sheet = workbook.create_sheet("DCF")
        sheet.append(["Metric", "Value"])
        for key, value in dcf_result.items():
            if key == "warnings":
                value = "; ".join(str(item) for item in value) if value else ""
            sheet.append([key.replace("_", " ").title(), _scalar(value)])
        style_header(sheet)
        autosize(sheet)

    sensitivity = body.get("sensitivity") or {}
    if sensitivity and sensitivity.get("table"):
        sheet = workbook.create_sheet("Sensitivity")
        terminal = sensitivity.get("terminal_values") or []
        sheet.append(["WACC \\ Terminal"] + [value for value in terminal])
        for row_index, wacc_value in enumerate(sensitivity.get("wacc_values") or []):
            values = sensitivity["table"][row_index] if row_index < len(sensitivity["table"]) else []
            sheet.append([wacc_value] + list(values))
        style_header(sheet)
        autosize(sheet)

    report = body.get("report") or {}
    if report:
        sheet = workbook.create_sheet("Report")
        sheet.append(["Section", "Content"])
        for key, value in report.items():
            if isinstance(value, (list, tuple)):
                rendered = "\n".join(f"- {item}" for item in value)
            elif isinstance(value, dict):
                rendered = str(value)
            else:
                rendered = str(value)
            sheet.append([key.replace("_", " ").title(), rendered])
            sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        style_header(sheet)
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 100

    quality = body.get("data_quality") or {}
    sheet = workbook.create_sheet("Sources")
    sheet.append(["APEX Terminal export"])
    sheet.append(["Generated", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    sheet.append(["Ticker", str(body.get("ticker") or "")])
    sheet.append(["Data source", "Yahoo Finance via yfinance"])
    for alert in quality.get("alerts") or []:
        sheet.append([str(alert.get("level", "info")).title(), str(alert.get("message", ""))])
    sheet["A1"].font = bold
    autosize(sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename(body, "model.xlsx")}"'},
    )


@router.post("/pdf")
def export_pdf(body: dict[str, Any]):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
        from reportlab.platypus import (
            HRFlowable,
            Image,
            KeepTogether,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"PDF export unavailable: {error}")

    title = str(body.get("title") or "APEX Terminal Research Report")
    ticker_text = _sanitize(str(body.get("ticker") or ""))
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    company = body.get("company") if isinstance(body.get("company"), dict) else {}
    actuals = body.get("actuals") if isinstance(body.get("actuals"), dict) else {}
    price_summary = body.get("price_summary") if isinstance(body.get("price_summary"), dict) else {}
    currency = str(company.get("currency") or "")
    company_name = _sanitize(str(company.get("name") or "")).strip()
    base_fy = _fy_year(company.get("last_data_date"))
    rating = body.get("rating") if isinstance(body.get("rating"), dict) else {}

    # Preferred path: use the deterministic note payload returned with the AI note.
    research_payload = body.get("research_payload")
    if not isinstance(research_payload, dict):
        research_payload = body.get("payload") if isinstance(body.get("payload"), dict) else None
    if isinstance(research_payload, dict):
        payload_meta = research_payload.get("meta") if isinstance(research_payload.get("meta"), dict) else {}
        if payload_meta:
            company = payload_meta
        hist_metrics = (research_payload.get("historical") or {}).get("metrics") or []
        if hist_metrics and isinstance(hist_metrics[0], dict):
            actuals = hist_metrics[0]
            base_fy = _fy_year(hist_metrics[0].get("period_end")) or base_fy
            if not company.get("last_data_date"):
                company = {**company, "last_data_date": hist_metrics[0].get("period_end")}
        payload_price = research_payload.get("price") if isinstance(research_payload.get("price"), dict) else {}
        if payload_price.get("current") is not None:
            price_summary = {
                **(body.get("price_summary") if isinstance(body.get("price_summary"), dict) else {}),
                **{k: payload_price.get(k) for k in ("current", "price_date", "change_pct", "low_52w", "high_52w")},
                "last_close": payload_price.get("current"),
            }
        else:
            price_summary = body.get("price_summary") if isinstance(body.get("price_summary"), dict) else {}
        payload_forecast = research_payload.get("forecast") if isinstance(research_payload.get("forecast"), dict) else {}
        if payload_forecast:
            forecast = payload_forecast
        payload_valuation = research_payload.get("valuation") if isinstance(research_payload.get("valuation"), dict) else {}
        payload_dcf = payload_valuation.get("dcf") if isinstance(payload_valuation.get("dcf"), dict) else {}
        if payload_dcf:
            dcf_result = payload_dcf
        payload_rating = research_payload.get("rating") if isinstance(research_payload.get("rating"), dict) else {}
        if payload_rating:
            rating = payload_rating

    report = body.get("report") if isinstance(body.get("report"), dict) else {}
    if isinstance(research_payload, dict):
        dcf_result = payload_dcf if isinstance(payload_dcf, dict) and payload_dcf else dcf_result
        forecast = payload_forecast if isinstance(payload_forecast, dict) and payload_forecast else forecast
    else:
        dcf_result = body.get("dcf") if isinstance(body.get("dcf"), dict) else dcf_result
        forecast = body.get("forecast") if isinstance(body.get("forecast"), dict) else forecast
    periods = [p for p in (forecast.get("periods") or []) if isinstance(p, dict)]
    assumptions_raw = forecast.get("assumptions") or body.get("assumptions") if isinstance(research_payload, dict) else (body.get("assumptions") or forecast.get("assumptions"))
    if isinstance(assumptions_raw, dict) and isinstance(assumptions_raw.get("values"), dict):
        assumptions = assumptions_raw.get("values")
    else:
        assumptions = assumptions_raw if isinstance(assumptions_raw, dict) else {}
    currency = currency or str(company.get("currency") or "")

    letter_w = letter[0]
    left_margin = 0.72 * inch
    right_margin = 0.72 * inch
    top_margin = 0.68 * inch
    bottom_margin = 0.78 * inch
    content_width = letter_w - left_margin - right_margin

    kicker_style = ParagraphStyle(
        "kicker",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(_ACCENT),
        spaceAfter=3,
    )
    h1_style = ParagraphStyle(
        "h1",
        fontName="Helvetica-Bold",
        fontSize=16.5,
        leading=19.5,
        textColor=colors.HexColor(_INK),
        spaceAfter=1,
    )
    meta_style = ParagraphStyle(
        "meta",
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor(_MUTED),
    )
    meta_right_style = ParagraphStyle(
        "meta_right",
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor(_MUTED),
        alignment=TA_RIGHT,
    )
    h2_style = ParagraphStyle(
        "h2",
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=14,
        textColor=colors.HexColor(_INK),
    )
    body_style = ParagraphStyle(
        "body",
        fontName="Helvetica",
        fontSize=9.2,
        leading=13.2,
        textColor=colors.HexColor(_BODY),
        spaceAfter=5.5,
    )
    bullet_style = ParagraphStyle(
        "bullet",
        fontName="Helvetica",
        fontSize=9.2,
        leading=13.2,
        textColor=colors.HexColor(_BODY),
        leftIndent=16,
        bulletIndent=4,
        bulletFontName="Helvetica",
        bulletFontSize=9,
        bulletColor=colors.HexColor(_ACCENT),
        spaceAfter=3.5,
    )
    small_style = ParagraphStyle(
        "small",
        fontName="Helvetica",
        fontSize=7.6,
        leading=10.6,
        textColor=colors.HexColor(_MUTED),
        spaceAfter=3,
    )
    kpi_style = ParagraphStyle(
        "kpi",
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor(_INK),
    )
    box_title_style = ParagraphStyle(
        "box_title",
        fontName="Helvetica-Bold",
        fontSize=8.6,
        leading=11.5,
        textColor=colors.HexColor(_INK),
        spaceAfter=2,
    )
    box_title_warn_style = ParagraphStyle(
        "box_title_warn",
        fontName="Helvetica-Bold",
        fontSize=8.6,
        leading=11.5,
        textColor=colors.HexColor(_WARN_TEXT),
        spaceAfter=2,
    )
    box_line_style = ParagraphStyle(
        "box_line",
        fontName="Helvetica",
        fontSize=8.3,
        leading=11.4,
        textColor=colors.HexColor(_BODY),
        leftIndent=14,
        bulletIndent=2,
        bulletFontName="Helvetica",
        bulletFontSize=8,
        spaceAfter=1.5,
    )
    cell_label_style = ParagraphStyle(
        "cell_label",
        fontName="Helvetica-Bold",
        fontSize=8.3,
        leading=10.4,
        textColor=colors.HexColor(_INK),
    )
    cell_text_style = ParagraphStyle(
        "cell_text",
        fontName="Helvetica",
        fontSize=8.0,
        leading=10.4,
        textColor=colors.HexColor(_BODY),
    )

    def note_box(title: str | None, lines: list[str], *, warn: bool) -> Table:
        title_style = box_title_warn_style if warn else box_title_style
        flowables: list[Any] = []
        if title:
            flowables.append(Paragraph(escape(title), title_style))
        for line in lines:
            flowables.append(Paragraph(escape(line), box_line_style, bulletText="•"))
        box = Table([[flowables]], colWidths=[content_width])
        box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(_WARN_BG if warn else "#F1F5F9")),
                    ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(_WARN_BORDER if warn else _RULE)),
                    ("LEFTPADDING", (0, 0), (-1, -1), 9),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        return box

    def data_table_style() -> TableStyle:
        return TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_INK)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7.6),
                ("FONTSIZE", (0, 1), (-1, -1), 8.4),
                ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor(_BODY)),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor(_RULE)),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(_RULE)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(_PANEL)]),
            ]
        )

    story: list[Any] = []

    story.append(Paragraph("APEX TERMINAL · EQUITY RESEARCH NOTE", kicker_style))
    story.append(Paragraph(escape(company_name or _sanitize(title)), h1_style))

    exchange = _sanitize(str(company.get("exchange") or ""))
    meta_left_top = "  ·  ".join(part for part in (ticker_text, exchange, currency.upper()) if part) or "APEX Terminal"
    sector = _sanitize(str(company.get("sector") or ""))
    industry = _sanitize(str(company.get("industry") or ""))
    meta_left_bottom = "  ·  ".join(part for part in (sector, industry) if part)
    last_date = _sanitize(str(company.get("last_data_date") or ""))
    meta_right_bottom = f"Data as of {last_date}" if last_date else ""

    meta_rows: list[list[Any]] = [
        [
            Paragraph(escape(meta_left_top), meta_style),
            Paragraph(escape(f"Generated {generated}"), meta_right_style),
        ]
    ]
    if meta_left_bottom or meta_right_bottom:
        meta_rows.append(
            [
                Paragraph(escape(meta_left_bottom), meta_style),
                Paragraph(escape(meta_right_bottom), meta_right_style),
            ]
        )
    meta_table = Table(meta_rows, colWidths=[content_width * 0.6, content_width * 0.4])
    meta_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    story.append(meta_table)
    story.append(
        HRFlowable(
            width="24%",
            thickness=1.6,
            color=colors.HexColor(_ACCENT),
            spaceBefore=7,
            spaceAfter=1,
            hAlign="LEFT",
        )
    )

    recommendation_body = {
        "rating": rating,
        "dcf": dcf_result,
        "price_summary": price_summary,
        "company": company,
    }
    recommendation_rows = _recommendation_rows(recommendation_body)
    if recommendation_rows:
        rec_table = Table(
            [[
                Paragraph(escape(_sanitize(label)), ParagraphStyle("rec_label", fontName="Helvetica-Bold", fontSize=8.3, leading=10.4, textColor=colors.HexColor(_MUTED))),
                Paragraph(
                    f'<font size="10.2" color="{_NEG if str(value).startswith("-") else _INK}"><b>{escape(_sanitize(value))}</b></font><br/>'
                    f'<font size="6.6" color="{_MUTED}">{escape(_sanitize(sub))}</font>',
                    kpi_style,
                ),
            ] for label, value, sub in recommendation_rows],
            colWidths=[content_width * 0.30, content_width * 0.70],
            hAlign="LEFT",
        )
        rec_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(_PANEL)),
                    ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(_RULE)),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 9),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(Spacer(1, 8))
        story.append(rec_table)

    kpis = _kpi_rows(body)
    if len(kpis) >= 2:
        kpi_cells: list[Any] = []
        for label, value, sub in kpis:
            value_color = _NEG if value.startswith("-") else _INK
            kpi_cells.append(
                Paragraph(
                    f'<font size="6.6" color="{_MUTED}">{escape(label.upper())}</font><br/>'
                    f'<font size="10.5" color="{value_color}"><b>{escape(value)}</b></font><br/>'
                    f'<font size="6.6" color="{_MUTED}">{escape(_sanitize(sub))}</font>',
                    kpi_style,
                )
            )
        kpi_row_list = [kpi_cells[i : i + 4] for i in range(0, len(kpi_cells), 4)]
        if len(kpi_row_list[-1]) < 4:
            kpi_row_list[-1] += [""] * (4 - len(kpi_row_list[-1]))
        kpi_table = Table(kpi_row_list, colWidths=[content_width / 4] * 4)
        kpi_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(_PANEL)),
                    ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(_RULE)),
                    ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor(_RULE)),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(Spacer(1, 10))
        story.append(kpi_table)

    price_history = []
    if isinstance(research_payload, dict):
        price_history = research_payload.get("price", {}).get("series") or []
    else:
        price_history = body.get("price_history") or []
    if price_history:
        target_for_chart = rating.get("target_price")
        if target_for_chart is None:
            target_for_chart = dcf_result.get("value_per_share")
        chart_png = _price_chart_png(price_history, target_for_chart, currency)
        if chart_png:
            image_width = content_width * 0.78
            image = Image(io.BytesIO(chart_png), width=image_width, height=image_width * 2.0 / 7.0)
            image.hAlign = "LEFT"
            story.append(Spacer(1, 9))
            story.append(image)
            story.append(Paragraph(escape(_sanitize("Exhibit 1 - Daily close history; dashed line = deterministic model target (E).")), small_style))

    def append_section(heading_text: str, flowables: list[Any]) -> None:
        rule = HRFlowable(
            width="100%",
            thickness=0.5,
            color=colors.HexColor(_RULE),
            spaceBefore=12,
            spaceAfter=4,
            hAlign="LEFT",
        )
        heading = Paragraph(escape(_sanitize(heading_text)), h2_style)
        if flowables:
            story.append(KeepTogether([rule, heading, flowables[0]]))
            story.extend(flowables[1:])
        else:
            story.append(KeepTogether([rule, heading]))

    ordered_keys = [key for key in _SECTION_ORDER if key in report]
    ordered_keys += [key for key in report if key not in _SECTION_ORDER and key != "figure_warnings"]

    for key in ordered_keys:
        content = report.get(key)
        if not content:
            continue
        section_flowables: list[Any] = []
        if isinstance(content, (list, tuple)):
            for item in content:
                section_flowables.append(Paragraph(escape(_sanitize(item)), bullet_style, bulletText="•"))
        elif isinstance(content, dict):
            for sub_key, sub_value in content.items():
                section_flowables.append(
                    Paragraph(
                        f"<b>{escape(_sanitize(sub_key))}</b>: {escape(_sanitize(sub_value))}",
                        body_style,
                    )
                )
        else:
            section_flowables.append(Paragraph(escape(_sanitize(content)), body_style))
        append_section(_humanize(str(key)), section_flowables)

    figure_warnings = report.get("figure_warnings")
    if isinstance(figure_warnings, list) and figure_warnings:
        append_section(
            "Figure warnings",
            [note_box(None, [_sanitize(str(w)) for w in figure_warnings], warn=True)],
        )

    dcf_rows: list[list[Any]] = []
    dcf_warnings: list[str] = []
    for key, value in dcf_result.items():
        if key == "warnings":
            if isinstance(value, (list, tuple)):
                dcf_warnings = [_sanitize(str(w)) for w in value]
            continue
        dcf_rows.append(
            [
                Paragraph(escape(_sanitize(_humanize(str(key)))), cell_label_style),
                _fmt_dcf_value(str(key), value, currency),
            ]
        )
    if dcf_rows:
        dcf_table = Table(
            [["Metric", "Value"]] + dcf_rows,
            colWidths=[content_width * 0.6, content_width * 0.4],
        )
        dcf_table.setStyle(data_table_style())
        append_section("Valuation Detail (DCF)", [dcf_table])
        if dcf_warnings:
            story.append(Spacer(1, 6))
            story.append(note_box("Model warnings", dcf_warnings, warn=True))

    assume_order = [
        "base_period",
        "revenue",
        "gross_margin",
        "opex_ratio",
        "tax_rate",
        "capex_pct",
        "da_pct",
        "interest_rate",
        "dividend_payout",
        "shares",
    ]
    assume_labels = {
        "base_period": "Base period",
        "revenue": "Base revenue",
        "gross_margin": "Gross margin",
        "opex_ratio": "Opex / revenue",
        "tax_rate": "Tax rate",
        "capex_pct": "Capex / revenue",
        "da_pct": "D&A / revenue",
        "interest_rate": "Interest rate",
        "dividend_payout": "Dividend payout",
        "shares": "Shares outstanding",
    }
    assume_percent = {"gross_margin", "opex_ratio", "tax_rate", "capex_pct", "da_pct", "interest_rate", "dividend_payout"}
    assumption_items: list[tuple[str, str]] = []
    for key in assume_order:
        if key not in assumptions:
            continue
        raw = assumptions.get(key)
        label = assume_labels.get(key, _humanize(key))
        if isinstance(raw, (list, tuple)):
            if len(raw) == 1:
                raw = raw[0]
            else:
                formatted = "per-year schedule"
                assumption_items.append((label, formatted))
                continue
        if key in assume_percent:
            formatted = _fmt_pct(raw) or "—"
        elif key in {"revenue", "shares"}:
            formatted = _fmt_compact(raw)
        else:
            formatted = _sanitize(raw)
        assumption_items.append((label, formatted))
    base_revenue = _f(assumptions.get("revenue"))
    first_revenue = _f(periods[0].get("revenue")) if periods else None
    if base_revenue not in (None, 0) and first_revenue is not None:
        assumption_items.append(("Revenue growth (model)", _fmt_pct(first_revenue / base_revenue - 1.0, signed=True) or "—"))

    if assumption_items:
        pair_rows: list[list[Any]] = []
        for index in range(0, len(assumption_items), 2):
            left_label, left_value = assumption_items[index]
            right = assumption_items[index + 1] if index + 1 < len(assumption_items) else ("", "")
            pair_rows.append(
                [
                    Paragraph(escape(left_label), cell_label_style) if left_label else "",
                    left_value,
                    Paragraph(escape(right[0]), cell_label_style) if right[0] else "",
                    right[1],
                ]
            )
        assume_table = Table(
            pair_rows,
            colWidths=[content_width * 0.29, content_width * 0.21, content_width * 0.29, content_width * 0.21],
        )
        assume_table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 8.4),
                    ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(_BODY)),
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("ALIGN", (3, 0), (3, -1), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4.5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor(_RULE)),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(_RULE)),
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor(_PANEL)]),
                ]
            )
        )
        append_section("Model assumptions", [assume_table])

    if periods:
        forecast_labels = [str(label) for label in (forecast.get("labels") or []) if label]
        show_actual = _has_actuals(actuals)
        header_row: list[Any] = ["Line item"]
        if show_actual:
            header_row.append(f"FY{base_fy} (A)" if base_fy else "Latest (A)")
        for index, period in enumerate(periods):
            if index < len(forecast_labels):
                header_row.append(f"{_sanitize(forecast_labels[index])} (P)")
                continue
            year_value = _f(period.get("year"))
            year_int = int(year_value) if year_value else index + 1
            if base_fy and 1 <= year_int <= 12:
                fiscal_year = base_fy + year_int
            elif base_fy and base_fy <= year_int <= base_fy + 30:
                fiscal_year = year_int
            elif base_fy:
                fiscal_year = base_fy + index + 1
            else:
                fiscal_year = None
            header_row.append(f"FY{fiscal_year} (P)" if fiscal_year else f"Y{year_int} (P)")

        data_rows: list[list[Any]] = [header_row]
        for label, key in _FORECAST_LINES:
            row: list[Any] = [Paragraph(escape(label), cell_label_style)]
            if show_actual:
                row.append(_fmt_compact(actuals.get(key)))
            for period in periods:
                row.append(_fmt_compact(period.get(key)))
            data_rows.append(row)

        data_cols = len(header_row) - 1
        col_widths = [content_width * 0.28] + [(content_width * 0.72) / data_cols] * data_cols
        forecast_table = Table(data_rows, colWidths=col_widths, repeatRows=1)
        style_commands = data_table_style().getCommands()
        if show_actual:
            style_commands.append(("BACKGROUND", (1, 1), (1, -1), colors.HexColor("#EFF6FF")))
        for row_index in range(1, len(data_rows)):
            for col_index in range(1, len(header_row)):
                cell_value = data_rows[row_index][col_index]
                if isinstance(cell_value, str) and cell_value.startswith("-"):
                    style_commands.append(
                        ("TEXTCOLOR", (col_index, row_index), (col_index, row_index), colors.HexColor(_NEG))
                    )
        forecast_table.setStyle(TableStyle(style_commands))
        append_section("Forecast Summary", [forecast_table])

        currency_note = currency.upper() or "local currency"
        story.append(Paragraph(escape(f"Amounts in {currency_note}; bn = billions, m = millions."), small_style))
        legend = "(P) deterministic model projection, not company guidance"
        if show_actual:
            actual_label = f" for FY{base_fy}" if base_fy else ""
            legend = f"(A) reported actual{actual_label} · {legend}"
        story.append(Paragraph(escape(legend), small_style))
        if show_actual:
            story.append(
                Paragraph(
                    "Note: model FCF = operating cash flow - capex. Reported cash-flow statements may use the opposite capex sign.",
                    small_style,
                )
            )

    quality = body.get("data_quality") if isinstance(body.get("data_quality"), dict) else {}
    alert_lines: list[str] = []
    for alert in quality.get("alerts") or []:
        if isinstance(alert, dict):
            level = str(alert.get("level") or "info").upper()
            message = _sanitize(str(alert.get("message") or ""))
            alert_lines.append(f"{level}: {message}" if message else level)
        else:
            alert_lines.append(_sanitize(str(alert)))
    if alert_lines:
        append_section("Data quality", [note_box(None, alert_lines, warn=False)])

    if isinstance(research_payload, dict):
        payload_sources = research_payload.get("sources") or []
        if payload_sources:
            source_flowables: list[Any] = [
                Table(
                    [["Ref", "Source", "Scope", "Retrieved"]] + [
                        [
                            Paragraph(escape(_sanitize(str(source.get("ref") or ""))), cell_text_style),
                            Paragraph(escape(_sanitize(str(source.get("label") or ""))), cell_text_style),
                            Paragraph(escape(_sanitize(str(source.get("where") or ""))), cell_text_style),
                            Paragraph(escape(_sanitize(str(source.get("retrieved_at") or "—"))), cell_text_style),
                        ]
                        for source in payload_sources[:8]
                    ],
                    colWidths=[content_width * 0.09, content_width * 0.30, content_width * 0.43, content_width * 0.18],
                    repeatRows=1,
                )
            ]
            source_flowables[0].setStyle(data_table_style())
            append_section("Provenance register", source_flowables)

    if isinstance(research_payload, dict):
        scenarios = research_payload.get("scenarios") if isinstance(research_payload.get("scenarios"), dict) else {}
        scenario_rows: list[list[Any]] = [["Scenario", "Value per share", "Upside", "Assumption"]]
        current_price_for_scenario = _f(price_summary.get("last_close") or price_summary.get("current"))
        for label in ("bear", "base", "bull"):
            scenario = scenarios.get(label) if isinstance(scenarios.get(label), dict) else {}
            value = _f(scenario.get("target_price"))
            if value is None:
                continue
            upside = value / current_price_for_scenario - 1 if current_price_for_scenario else None
            growth_delta = _f(scenario.get("growth_delta"))
            premise = "APEX scenario design" if label == "base" else (
                f"Revenue growth {_fmt_pct(growth_delta, signed=True) or 'adjusted'} and WACC {_fmt_pct(scenario.get('wacc')) or 'adjusted'}"
            )
            scenario_rows.append([
                Paragraph(escape(_sanitize(label.capitalize())), cell_label_style),
                _fmt_money(value, currency),
                _fmt_pct(upside, signed=True) or "—",
                Paragraph(escape(_sanitize(premise)), cell_text_style),
            ])
        if len(scenario_rows) > 1:
            scenario_table = Table(scenario_rows, colWidths=[content_width * 0.16, content_width * 0.22, content_width * 0.16, content_width * 0.46], repeatRows=1)
            scenario_table.setStyle(data_table_style())
            append_section("Deterministic scenarios", [scenario_table])
            story.append(Paragraph(escape(_sanitize("Bear/base/bull are model outputs; they are not company guidance.")), small_style))

    append_section(
        "Sources & Limitations",
        [
            Paragraph(
                escape(
                    _sanitize(
                        "Data: Yahoo Finance via yfinance. Forecasts are deterministic outputs of the APEX driver model. "
                        "AI text is assistive only and never supplies final figures. Missing data is shown as a dash, not zero. "
                        "This document is not investment advice."
                    )
                ),
                small_style,
            )
        ],
    )

    footer_left = f"APEX Terminal  ·  {ticker_text}" if ticker_text else "APEX Terminal"

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            super().__init__(*args, **kwargs)
            self._page_states: list[dict[str, Any]] = []

        def showPage(self) -> None:
            self._page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self) -> None:
            total = len(self._page_states)
            for page_state in self._page_states:
                self.__dict__.update(page_state)
                self._draw_footer(total)
                super().showPage()
            super().save()

        def _draw_footer(self, total: int) -> None:
            self.saveState()
            self.setStrokeColor(colors.HexColor(_RULE))
            self.setLineWidth(0.5)
            self.line(left_margin, bottom_margin - 18, letter_w - right_margin, bottom_margin - 18)
            self.setFont("Helvetica", 7.2)
            self.setFillColor(colors.HexColor(_MUTED))
            self.drawString(left_margin, bottom_margin - 30, footer_left)
            self.drawRightString(letter_w - right_margin, bottom_margin - 30, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
        title=title,
        author="APEX Terminal",
        subject=f"{ticker_text} research note",
    )
    document.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)

    qa_headers: dict[str, str] = {}
    if isinstance(research_payload, dict):
        from ..services.pdf_qa import run_note_pdf_qa

        qa_issues = run_note_pdf_qa(buffer.getvalue(), {
            "rating": rating,
            "dcf": dcf_result,
            "forecast": forecast,
        })
        blocking = [issue for issue in qa_issues if issue.get("level") == "error"]
        if blocking:
            raise HTTPException(status_code=422, detail={"message": "Concise-note QA blocked publication.", "issues": blocking})
        qa_headers["X-APEX-QC-Status"] = "passed" if any(issue.get("check") == "pdf.passed" for issue in qa_issues) else "warnings"
        qa_headers["X-APEX-QC-Issues"] = "; ".join(f'{issue.get("check")}:{issue.get("level")}' for issue in qa_issues if issue.get("level") != "info")

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_filename(body, "report.pdf")}"', **qa_headers},
    )
